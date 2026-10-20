#!/usr/bin/env python3
"""Bundle the review's working files into one Excel evidence workbook.

Usage:
    python3 scripts/export_workbook.py litrev-<slug>/ [-o evidence_workbook.xlsx]

One sheet per file that exists: Search log, Records (de-duplicated), Duplicates,
Screening, Verification, Full text, Study cards, Weaknesses, Extraction, Estimates,
Appraisal, Seeds, PRISMA counts, plus a README sheet listing what each sheet holds and
the file it came from. Long text cells (cards, weaknesses) wrap.
Headers are frozen and filterable. Requires openpyxl (pip install openpyxl); without
it the CSV files themselves are the evidence appendix.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from litrev_common import FILES, read_csv  # noqa: E402

SHEETS = [
    ("Search log", "search_log", "Every query: date, source, string, filters, hits, retrieved"),
    ("Records", "records_dedup", "Unique records after de-duplication"),
    ("Duplicates", "duplicates", "Removed duplicates and working-paper/journal pairs"),
    ("Screening", "screening", "Title/abstract and full-text decisions with exclusion codes"),
    ("Verification", "verification", "Verified metadata, link, check method and date, status"),
    ("Full text", "fulltext", "Legal open-access copies found (Unpaywall, OpenAlex, Semantic Scholar)"),
    ("Study cards", "cards", "Three-pass summaries: what each study did and what it left out"),
    ("Weaknesses", "weaknesses", "Weaknesses by category and severity, with evidence location and basis"),
    ("Extraction", "extraction", "Coded study characteristics and findings with locations"),
    ("Estimates", "estimates", "Comparable estimates used in the meta-analysis"),
    ("Appraisal", "appraisal", "Credibility scores and venue information with sources"),
    ("Seeds", "seeds", "Known core papers used to test search recall"),
]


WRAP_SHEETS = {"Study cards", "Weaknesses"}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("-o", "--output", default="")
    args = ap.parse_args(argv)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl is not installed (pip install openpyxl); deliver the CSV files instead.")
        return 2

    folder = Path(args.folder)
    out = Path(args.output) if args.output else folder / "evidence_workbook.xlsx"
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["Sheet", "Source file", "Contents", "Rows"])
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3B57")

    def style_header(ws):
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"

    written = 0
    for title, key, desc in SHEETS:
        path = folder / FILES[key]
        if key == "records_dedup" and not path.exists():
            path = folder / FILES["records"]
        if not path.exists():
            continue
        rows = read_csv(path)
        if not rows:
            continue
        ws = wb.create_sheet(title)
        header = list(rows[0].keys())
        ws.append(header)
        for r in rows:
            ws.append([r.get(h, "") for h in header])
        style_header(ws)
        ws.auto_filter.ref = ws.dimensions
        for j, h in enumerate(header, 1):
            width = max(len(str(h)), *(min(len(str(r.get(h, ""))), 60) for r in rows[:200]))
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 62)
        if title in WRAP_SHEETS:
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        readme.append([title, path.name, desc, len(rows)])
        written += 1

    counts_path = folder / FILES["prisma_counts"]
    if counts_path.exists():
        counts = json.loads(counts_path.read_text(encoding="utf-8"))
        ws = wb.create_sheet("PRISMA counts")
        ws.append(["Item", "Value"])
        for k, v in counts.items():
            ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
        style_header(ws)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 60
        readme.append(["PRISMA counts", counts_path.name, "PRISMA 2020 flow counts", len(counts)])
        written += 1

    style_header(readme)
    for col, w in zip("ABCD", (16, 26, 70, 8)):
        readme.column_dimensions[col].width = w
    if not written:
        print("No working files with rows found - nothing to export.")
        return 1
    wb.save(out)
    print(f"Wrote {out} ({written} sheets + README)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
