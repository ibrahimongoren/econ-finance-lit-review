#!/usr/bin/env python3
"""Build the illustrated report (DOCX and/or PDF) from 10_report.md, figures, study cards and references.

Usage:
    python3 scripts/build_report.py litrev-<slug>/ [--format docx|pdf|both] [--lang tr|en]
        [--title "..."] [--subtitle "..."] [--author "..."] [--toc] [--no-figures] [--soffice PATH]

Write the narrative in 10_report.md (init_review.py creates a skeleton: answer first, then scope, search,
shape of the literature, findings, weaknesses, gaps, limitations) and keep placeholders where generated
material belongs; every placeholder is optional:
    {{FIGURE:pipeline}} {{FIGURE:prisma}} {{FIGURE:years}} {{FIGURE:map}} {{FIGURE:weaknesses}}
    {{FIGURE:regions}} {{FIGURE:sources}}     figures with numbered captions
    {{CARDS}}             every study card: three passes, what it did / left out, weaknesses
    {{WEAKNESS_TABLE}}    weaknesses by category and severity
    {{REFERENCES}}        APA 7 reference list generated from 05_verification.csv
    {{TALLY}}             verification tally line
If {{CARDS}} or {{REFERENCES}} is missing, those sections are appended at the end. HTML comments
(<!-- ... -->) are dropped. The cover page gets a key-facts table (window, studies, sources, search dates,
verification); --toc adds a contents list with links to every section.

Pipeline: make_figures.py -> 10_report_full.md -> pandoc with a styled reference.docx -> 10_report.docx
-> LibreOffice (soffice) -> 10_report.pdf. Without pandoc a built-in python-docx renderer is used; without
LibreOffice the DOCX is delivered and the PDF step is reported as skipped (export it from Word instead).
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from litrev_common import (FILES, PRINTABLE_STATUSES, SEVERITIES, WEAKNESS_CATEGORIES,  # noqa: E402
                           WEAKNESS_LABELS, read_csv, window_from_protocol)

INK, INK2, HEAD1, HEAD2, LINK = "1D1D1B", "52514E", "1F3B57", "2E5A85", "2A78D6"
FIG_NAMES = ["pipeline", "prisma", "years", "map", "weaknesses", "regions", "sources"]
CAPTION = {
    "en": {"pipeline": "Review pipeline and the file each stage leaves behind",
           "prisma": "PRISMA 2020 flow diagram", "years": "Included studies by publication year and method family",
           "map": "Evidence map: themes × method families (number of studies)",
           "weaknesses": "Weaknesses identified across included studies, by category and severity",
           "regions": "Included studies by country or region", "sources": "Records identified per source",
           "fig": "Figure", "cards": "Study cards", "refs": "References", "weak_tbl": "Weaknesses by category",
           "cat": "Category", "high": "High", "medium": "Medium", "low": "Low", "total": "Total",
           "studies": "Studies", "tally": "Verification", "made": "Prepared with econ-finance-lit-review",
           "contents": "Contents", "window": "Literature window", "included": "Studies included",
           "ft_ab": "{ft} read in full text, {ab} from the abstract only", "sources_searched": "Sources searched",
           "dates": "Search dates", "verified": "References verified",
           "tally_words": ("verified", "partial", "working paper / preprint", "not verified", "retracted",
                           "links checked"),
           "months": ["January", "February", "March", "April", "May", "June", "July", "August", "September",
                      "October", "November", "December"]},
    "tr": {"pipeline": "Tarama süreci ve her aşamanın bıraktığı dosya", "prisma": "PRISMA 2020 akış şeması",
           "years": "Yayın yılına ve yöntem ailesine göre dahil edilen çalışmalar",
           "map": "Kanıt haritası: temalar × yöntem aileleri (çalışma sayısı)",
           "weaknesses": "Dahil edilen çalışmalarda tespit edilen zayıf yönler (kategori ve şiddet)",
           "regions": "Ülke veya bölgeye göre dahil edilen çalışmalar", "sources": "Kaynağa göre tanımlanan kayıtlar",
           "fig": "Şekil", "cards": "Çalışma kartları", "refs": "Kaynakça", "weak_tbl": "Kategoriye göre zayıf yönler",
           "cat": "Kategori", "high": "Yüksek", "medium": "Orta", "low": "Düşük", "total": "Toplam",
           "studies": "Çalışma", "tally": "Doğrulama", "made": "econ-finance-lit-review ile hazırlanmıştır",
           "contents": "İçindekiler", "window": "Literatür penceresi", "included": "Dahil edilen çalışma",
           "ft_ab": "{ft} tam metinden, {ab} yalnızca özetten okundu", "sources_searched": "Taranan kaynaklar",
           "dates": "Tarama tarihleri", "verified": "Doğrulanan kaynak",
           "tally_words": ("doğrulandı", "kısmen doğrulandı", "çalışma tebliği / ön baskı", "doğrulanamadı",
                           "geri çekilmiş", "bağlantı kontrolü"),
           "months": ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim",
                      "Kasım", "Aralık"]},
}
PAGE_BREAK = "```{=openxml}\n<w:p><w:r><w:br w:type=\"page\"/></w:r></w:p>\n```"
SOURCE_NAMES = {"openalex": "OpenAlex", "crossref": "Crossref", "semanticscholar": "Semantic Scholar",
                "arxiv": "arXiv", "scopus": "Scopus", "sciencedirect": "ScienceDirect",
                "webofscience": "Web of Science", "googlescholar-serpapi": "Google Scholar (SerpApi)",
                "googlescholar": "Google Scholar", "publishorperish": "Publish or Perish", "elicit": "Elicit",
                "consensus": "Consensus", "zotero": "Zotero", "econbiz": "EconBiz", "repec": "RePEc/IDEAS",
                "nber": "NBER", "ssrn": "SSRN"}


# ---------------------------------------------------------------------------
# Generated blocks
# ---------------------------------------------------------------------------
def figure_path(folder: Path, name: str) -> Path:
    if name == "prisma":
        return folder / "04_prisma_flow.png"
    return folder / FILES["figures_dir"] / f"fig_{name}.png"


def nice_date(iso: str, lang: str) -> str:
    try:
        d = dt.date.fromisoformat(iso)
    except ValueError:
        return iso
    return f"{d.day} {CAPTION[lang]['months'][d.month - 1]} {d.year}"


def weakness_table(folder: Path, lang: str) -> str:
    path = folder / FILES["weaknesses"]
    rows = read_csv(path) if path.exists() else []
    if not rows:
        return ""
    c = CAPTION[lang]
    counts, studies = defaultdict(Counter), defaultdict(set)
    for w in rows:
        counts[w.get("category")][(w.get("severity") or "").lower()] += 1
        studies[w.get("category")].add(w.get("ref_id"))
    # a separator longer than 72 characters makes pandoc use these relative column widths
    lines = [f"| {c['cat']} | {c['high']} | {c['medium']} | {c['low']} | {c['total']} | {c['studies']} |",
             "|" + "-" * 40 + "|" + "|".join(["-" * 9 + ":"] * 5) + "|"]
    for cat in sorted((k for k in counts if k in WEAKNESS_CATEGORIES), key=lambda k: -sum(counts[k].values())):
        sev = counts[cat]
        lines.append(f"| {WEAKNESS_LABELS[lang][cat]} | " + " | ".join(str(sev.get(s, 0)) for s in SEVERITIES)
                     + f" | {sum(sev.values())} | {len(studies[cat])} |")
    return f"**{c['weak_tbl']}**\n\n" + "\n".join(lines)


def tally(folder: Path, lang: str) -> str:
    ver = read_csv(folder / FILES["verification"]) if (folder / FILES["verification"]).exists() else []
    st = Counter((r.get("status") or "").upper() for r in ver)
    checked = sum(1 for r in ver if (r.get("status") or "").upper() in PRINTABLE_STATUSES
                  and (r.get("link_check_method") or "").strip())
    printable = sum(st[s] for s in PRINTABLE_STATUSES)
    w = CAPTION[lang]["tally_words"]
    return (f"**{CAPTION[lang]['tally']}:** {st['VERIFIED']} {w[0]} · {st['PARTIAL']} {w[1]} · "
            f"{st['PREPRINT']} {w[2]} · {st['NOT_VERIFIED']} {w[3]} · {st['RETRACTED']} {w[4]} · "
            f"{w[5]} {checked}/{printable}")


def pretty_source(raw: str, lang: str = "en") -> str:
    key = re.sub(r"[\s_]+", "", (raw or "").strip().lower())
    if key.startswith(("snowball", "citation")):
        return {"en": "Snowballing", "tr": "Kartopu taraması"}[lang]
    if key.startswith(("author", "seed", "user", "manual", "prior-review", "website")):
        return {"en": "Other methods", "tr": "Diğer yöntemler"}[lang]
    return SOURCE_NAMES.get(key, raw.strip())


def cover_facts(folder: Path, lang: str) -> str:
    c = CAPTION[lang]
    rows = []
    window = window_from_protocol(folder)
    if window:
        rows.append((c["window"], f"{window[0]}–{window[1]}"))
    cards_path = folder / FILES["cards"]
    cards = read_csv(cards_path) if cards_path.exists() else []
    if cards:
        levels = Counter((x.get("evidence_level") or "").upper().strip("[]") for x in cards)
        rows.append((c["included"], f"{len(cards)} (" + c["ft_ab"].format(ft=levels.get("FT", 0),
                                                                          ab=levels.get("AB", 0)) + ")"))
    log_path = folder / FILES["search_log"]
    log = read_csv(log_path) if log_path.exists() else []
    sources = []
    for r in log:
        name = pretty_source(r.get("source", ""), lang)
        if name and name not in sources:
            sources.append(name)
    if sources:
        rows.append((c["sources_searched"], ", ".join(sources)))
    dates = sorted(d for d in (r.get("date", "").strip()[:10] for r in log) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", d))
    if dates:
        rows.append((c["dates"], nice_date(dates[0], lang) + (f" – {nice_date(dates[-1], lang)}"
                                                              if dates[-1] != dates[0] else "")))
    ver_path = folder / FILES["verification"]
    if ver_path.exists():
        rows.append((c["verified"], tally(folder, lang).split(":** ", 1)[1]))
    if not rows:
        return ""
    lines = ["|   |   |", "|" + "-" * 30 + "|" + "-" * 60 + "|"]
    lines += [f"| **{k}** | {v} |" for k, v in rows]
    return "\n".join(lines)


def cards_md(folder: Path, lang: str) -> str:
    import study_cards as sc
    cards, weak, ver, _ = sc.load(folder)
    cards = [c for c in sc.ordered(cards, ver, "theme") if c["ref_id"] in ver]
    if not cards:
        return ""
    out, theme = [f"# {CAPTION[lang]['cards']}"], None
    for n, c in enumerate(cards, 1):
        if (c.get("theme") or "") != theme:
            theme = c.get("theme") or ""
            if theme:
                out.append(f"## {theme}")
        out.append(sc.render_card(n, c, weak.get(c["ref_id"], []), ver[c["ref_id"]], lang, report=True))
    return "\n\n".join(out)


def references_md(folder: Path, lang: str) -> str:
    import format_references as fr
    rows = read_csv(folder / FILES["verification"]) if (folder / FILES["verification"]).exists() else []
    entries = fr.reference_entries(rows, lang)
    if not entries:
        return ""
    return (f"# {CAPTION[lang]['refs']}\n\n::: {{custom-style=\"Bibliography\"}}\n" + "\n\n".join(entries)
            + "\n:::")


def add_contents(body: str, lang: str) -> str:
    """Give every level-1/2 heading an id and prepend a linked contents list (works in Word and PDF)."""
    out, toc, fence, k = [], [], False, 0
    for line in body.splitlines():
        if line.startswith("```"):
            fence = not fence
        m = None if fence else re.match(r"^(#{1,2}) (.+?)\s*(\{#([\w:.-]+)\})?\s*$", line)
        if m:
            k += 1
            ident = m.group(4) or f"sec-{k}"
            text = m.group(2).strip()
            line = f"{m.group(1)} {text} {{#{ident}}}"
            indent = "    " if len(m.group(1)) == 2 else ""
            toc.append(f"{indent}- [{re.sub(r'[*_]', '', text)}](#{ident})")
        out.append(line)
    if not toc:
        return body
    head = f"::: {{custom-style=\"TOC Heading\"}}\n{CAPTION[lang]['contents']}\n:::"
    return head + "\n\n" + "\n".join(toc) + "\n\n" + "\n".join(out)


def yaml_str(text: str) -> str:
    return '"' + (text or "").replace("\\", "\\\\").replace('"', '\\"') + '"'


def protocol_title(folder: Path) -> str:
    path = folder / FILES["protocol"]
    if path.exists():
        m = re.search(r"^# (?:Protocol|Protokol) — (.+)$", path.read_text(encoding="utf-8"), re.M)
        if m:
            return m.group(1).strip()
    return folder.name.replace("litrev-", "").replace("-", " ").title()


def assemble(folder: Path, lang: str, meta: dict, toc: bool = False) -> tuple[str, list[str]]:
    body = (folder / FILES["report_md"]).read_text(encoding="utf-8")
    body = re.sub(r"<!--.*?-->\s*", "", body, flags=re.S)
    notes, fig_no = [], 0

    def fig(m):
        nonlocal fig_no
        name = m.group(1).strip().lower()
        if name not in FIG_NAMES:
            notes.append(f"unknown figure '{name}' - use one of: {', '.join(FIG_NAMES)}")
            return ""
        path = figure_path(folder, name)
        if not path.exists():
            notes.append(f"figure '{name}' not found ({path.name}) - placeholder removed")
            return ""
        fig_no += 1
        rel = path.relative_to(folder).as_posix()
        return f"![{CAPTION[lang]['fig']} {fig_no}. {CAPTION[lang].get(name, name)}]({rel}){{width=16cm}}"

    body = re.sub(r"\{\{FIGURE:([\w-]+)\}\}", fig, body)
    body = body.replace("{{WEAKNESS_TABLE}}", weakness_table(folder, lang))
    body = body.replace("{{TALLY}}", tally(folder, lang))
    cards = cards_md(folder, lang)
    refs = references_md(folder, lang)
    body = body.replace("{{CARDS}}", cards) if "{{CARDS}}" in body else body + ("\n\n" + cards if cards else "")
    body = body.replace("{{REFERENCES}}", refs) if "{{REFERENCES}}" in body else body + ("\n\n" + refs if refs else "")
    body = re.sub(r"\n{3,}", "\n\n", body).strip()
    if toc:
        body = add_contents(body, lang)
    header = ["---", f"title: {yaml_str(meta['title'])}"]
    if meta.get("subtitle"):
        header.append(f"subtitle: {yaml_str(meta['subtitle'])}")
    if meta.get("author"):
        header.append(f"author: {yaml_str(meta['author'])}")
    header += [f"date: {yaml_str(meta['date'])}", f"lang: {'tr-TR' if lang == 'tr' else 'en-GB'}", "---"]
    facts = cover_facts(folder, lang)
    full = "\n".join(header) + "\n\n" + (facts + "\n\n" if facts else "") + PAGE_BREAK + "\n\n" + body + "\n"
    return full, notes


# ---------------------------------------------------------------------------
# DOCX styling
# ---------------------------------------------------------------------------
def find_style(doc, name: str):
    """Style by its displayed name. python-docx's doc.styles[...] maps 'Heading 1' to the internal
    'heading 1' and misses pandoc's capitalised names, so match on the name ourselves."""
    want = name.lower()
    return next((s for s in doc.styles if (s.name or "").lower() == want), None)


def style_document(doc, lang: str) -> None:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    if find_style(doc, "Bibliography") is None:
        doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH).base_style = find_style(doc, "Normal")
    if find_style(doc, "TOC Heading") is None:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH).base_style = find_style(doc, "Normal")

    def setf(name, size=None, bold=None, italic=None, color=None, before=None, after=None, font="Calibri",
             keep=False, align=None, hanging=None, spacing=None):
        st = find_style(doc, name)
        if st is None:
            return
        f = st.font
        f.name = font
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font)
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(theme_attr)) is not None:
                del rfonts.attrib[qn(theme_attr)]
        if size:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if italic is not None:
            f.italic = italic
        if color:
            f.color.rgb = RGBColor.from_string(color)
        pf = getattr(st, "paragraph_format", None)
        if pf is not None:
            if before is not None:
                pf.space_before = Pt(before)
            if after is not None:
                pf.space_after = Pt(after)
            if keep:
                pf.keep_with_next = True
            if align is not None:
                pf.alignment = align
            if hanging:
                pf.left_indent, pf.first_line_indent = Cm(hanging), Cm(-hanging)
            if spacing:
                pf.line_spacing = spacing

    for n in ("Normal", "Body Text", "First Paragraph"):
        setf(n, 10.5, color=INK, after=6, spacing=1.15)
    setf("Compact", 10.5, color=INK, after=2, spacing=1.1)
    setf("Title", 26, bold=True, color=HEAD1, before=120, after=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    setf("Subtitle", 14, color=HEAD2, after=24, align=WD_ALIGN_PARAGRAPH.LEFT)
    setf("Author", 11, color=INK2, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
    setf("Date", 11, color=INK2, after=30, align=WD_ALIGN_PARAGRAPH.LEFT)
    setf("Heading 1", 16, bold=True, color=HEAD1, before=20, after=8, keep=True)
    setf("Heading 2", 13, bold=True, color=HEAD2, before=14, after=6, keep=True)
    setf("Heading 3", 11, bold=True, color=HEAD1, before=12, after=3, keep=True)
    setf("Heading 4", 10.5, bold=True, color=INK2, before=8, after=3, keep=True)
    setf("Image Caption", 9, italic=True, color=INK2, after=12)
    setf("Table Caption", 9, italic=True, color=INK2, after=4)
    setf("Caption", 9, italic=True, color=INK2, after=12)
    setf("Block Text", 10, italic=True, color=INK2)
    setf("TOC Heading", 14, bold=True, color=HEAD1, before=6, after=8)
    setf("Bibliography", 10, color=INK, after=5, hanging=1.0, spacing=1.1)
    link_style = find_style(doc, "Hyperlink")
    if link_style is not None:
        link_style.font.color.rgb = RGBColor.from_string(LINK)
    table_style = find_style(doc, "Table")
    if table_style is not None:
        table_style.font.size = Pt(9.5)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.2)
        section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(CAPTION[lang]["made"] + "  ·  ")
        run.font.size, run.font.color.rgb = Pt(8), RGBColor.from_string("898781")
        r2 = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            r2._r.append(el)
        r2.font.size, r2.font.color.rgb = Pt(8), RGBColor.from_string("898781")


def make_reference_docx(path: Path, lang: str) -> bool:
    try:
        import docx
    except ImportError:
        return False
    res = subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"],
                         capture_output=True)
    if res.returncode != 0:
        return False
    d = docx.Document(str(path))
    style_document(d, lang)
    d.save(str(path))
    return True


# ---------------------------------------------------------------------------
# Fallback renderer (python-docx only)
# ---------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _inline(par, text: str):
    from docx.shared import RGBColor
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok).groups()
            if url.startswith("#"):
                par.add_run(label)
            else:
                r = par.add_run(f"{label} ({url})" if label != url else url)
                r.font.color.rgb = RGBColor.from_string(LINK)
        else:
            par.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def render_fallback(full_md: str, folder: Path, out: Path, lang: str) -> None:
    import docx
    from docx.shared import Cm
    doc = docx.Document()
    style_document(doc, lang)
    text = full_md
    meta = {}
    m = re.match(r"---\n(.*?)\n---\n", text, re.S)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip().strip('"').replace('\\"', '"')
        text = text[m.end():]
    for key, style in (("title", "Title"), ("subtitle", "Subtitle"), ("author", "Normal"), ("date", "Normal")):
        if meta.get(key):
            doc.add_paragraph(meta[key], style=style)
    lines = text.splitlines()
    i, para, div_style = 0, [], None

    def flush():
        if para:
            style = div_style if div_style and find_style(doc, div_style) is not None else None
            _inline(doc.add_paragraph(style=style), " ".join(s.strip() for s in para))
            para.clear()

    while i < len(lines):
        line = lines[i]
        if line.startswith("```{=openxml}"):
            flush()
            while i < len(lines) and lines[i].strip() != "```":
                i += 1
            doc.add_page_break()
        elif line.startswith(":::"):
            flush()
            sm = re.search(r'custom-style="([^"]+)"', line)
            div_style = sm.group(1) if sm else None
        elif re.match(r"^#{1,4} ", line):
            flush()
            level = len(line) - len(line.lstrip("#"))
            title = re.sub(r"\s*\{#[^}]*\}\s*$", "", line[level:].strip())
            title = re.sub(r"\*([^*]+)\*", r"\1", title)
            doc.add_heading(title, level=level)
        elif re.match(r"^\s*!\[", line):
            flush()
            im = re.match(r"^\s*!\[([^\]]*)\]\(([^)]+)\)", line)
            if im and (folder / im.group(2)).exists():
                doc.add_picture(str(folder / im.group(2)), width=Cm(16))
                doc.add_paragraph(im.group(1), style="Caption")
        elif re.match(r"^\s*[-*] ", line):
            flush()
            indent = len(line) - len(line.lstrip())
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            _inline(doc.add_paragraph(style=style), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            flush()
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", line))
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            i -= 1
            headerless = bool(rows) and not any(rows[0])
            rows = rows[1:] if headerless else rows
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _inline(cell.paragraphs[0], val)
                        if r == 0 and not headerless:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
        elif line.startswith(">"):
            flush()
            p = doc.add_paragraph(style="Normal")
            _inline(p, line.lstrip("> "))
            p.paragraph_format.left_indent = Cm(0.8)
        elif not line.strip():
            flush()
        else:
            para.append(line)
        i += 1
    flush()
    doc.save(str(out))


# ---------------------------------------------------------------------------
def find_soffice(explicit: str = "") -> str | None:
    candidates = [explicit] if explicit else []
    candidates += [shutil.which("soffice") or "", shutil.which("libreoffice") or "",
                   r"C:\Program Files\LibreOffice\program\soffice.exe",
                   "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    return next((c for c in candidates if c and Path(c).exists()), None)


def to_pdf(soffice: str, docx_path: Path) -> tuple[bool, str]:
    pdf = docx_path.with_suffix(".pdf")
    if pdf.exists():
        pdf.unlink()
    profile = Path(tempfile.mkdtemp(prefix="litrev-lo-"))
    try:
        res = subprocess.run([soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless",
                              "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                             capture_output=True, text=True, timeout=300)
    except (subprocess.TimeoutExpired, OSError) as e:
        return False, str(e)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    return pdf.exists(), (res.stderr or res.stdout).strip()[:300]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--format", choices=["docx", "pdf", "both"], default="both")
    ap.add_argument("--lang", choices=["tr", "en"], default="en")
    ap.add_argument("--title", default="", help="default: the title in 00_protocol.md")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--author", default="")
    ap.add_argument("--toc", action="store_true", help="add a linked contents list after the cover")
    ap.add_argument("--no-figures", action="store_true", help="reuse existing figures instead of redrawing")
    ap.add_argument("--fallback", action="store_true", help="force the python-docx renderer")
    ap.add_argument("--soffice", default="", help="path to LibreOffice soffice for the PDF step")
    a = ap.parse_args(argv)

    folder = Path(a.folder)
    if not (folder / FILES["report_md"]).exists():
        print(f"{FILES['report_md']} not found - write the narrative first (see references/reporting.md)")
        return 1
    if not a.no_figures:
        import make_figures
        make_figures.make(folder, a.lang)
    meta = {"title": a.title or protocol_title(folder), "subtitle": a.subtitle, "author": a.author,
            "date": nice_date(dt.date.today().isoformat(), a.lang)}
    full, notes = assemble(folder, a.lang, meta, toc=a.toc)
    full_path = folder / "10_report_full.md"
    full_path.write_text(full, encoding="utf-8")
    for n in notes:
        print("  · " + n)

    docx_path = folder / "10_report.docx"
    used = ""
    if shutil.which("pandoc") and not a.fallback:
        ref = folder / FILES["figures_dir"] / "reference.docx"
        ref.parent.mkdir(exist_ok=True)
        cmd = ["pandoc", full_path.name, "-o", docx_path.name,
               "-f", "markdown+pipe_tables+link_attributes+raw_attribute+implicit_figures+fenced_divs",
               "--resource-path", "."]
        if make_reference_docx(ref, a.lang):
            cmd += ["--reference-doc", str(ref.relative_to(folder))]
        res = subprocess.run(cmd, cwd=folder, capture_output=True, text=True)
        if res.returncode == 0:
            used = "pandoc"
        else:
            print("pandoc failed, using the built-in renderer: " + res.stderr.strip()[:300])
    if not used:
        try:
            render_fallback(full, folder, docx_path, a.lang)
            used = "python-docx"
        except ImportError:
            print("Neither pandoc nor python-docx is available - only 10_report_full.md was written.")
            return 1
    print(f"✓ {docx_path} ({used})")

    if a.format in ("pdf", "both"):
        soffice = find_soffice(a.soffice)
        if not soffice:
            print("· PDF skipped: LibreOffice (soffice) not found - open the DOCX and export it as PDF, "
                  "or pass --soffice PATH")
        else:
            ok, msg = to_pdf(soffice, docx_path)
            print(f"✓ {docx_path.with_suffix('.pdf')}" if ok else "· PDF conversion failed: " + msg)
    return 0


if __name__ == "__main__":
    sys.exit(main())
